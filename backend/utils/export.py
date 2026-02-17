"""
Export format generators for transcription jobs.
"""

import io
import json
from datetime import datetime

from config import SUPPORTED_LANGUAGES
from job_models import TranscriptionJob


def format_timestamp(seconds: float) -> str:
    """Format seconds to HH:MM:SS.mm"""
    hrs = int(seconds // 3600)
    mins = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    ms = int((seconds % 1) * 100)

    if hrs > 0:
        return f"{hrs:02d}:{mins:02d}:{secs:02d}.{ms:02d}"
    return f"{mins:02d}:{secs:02d}.{ms:02d}"


def format_srt_timestamp(seconds: float) -> str:
    """Format seconds to SRT timestamp format."""
    hrs = int(seconds // 3600)
    mins = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    ms = int((seconds % 1) * 1000)
    return f"{hrs:02d}:{mins:02d}:{secs:02d},{ms:03d}"


def generate_txt(job: TranscriptionJob, include_timestamps: bool = True, include_speakers: bool = True) -> str:
    """Generate plain text transcript."""
    lines = []

    for segment in job.segments:
        # Insert blank line for paragraph breaks
        if segment.get("paragraph_break") and lines:
            lines.append("")

        parts = []

        if include_timestamps:
            parts.append(f"[{format_timestamp(segment['start'])}]")

        if include_speakers and segment.get("speaker"):
            parts.append(f"{segment['speaker']}:")

        parts.append(segment["text"])
        lines.append(" ".join(parts))

    return "\n".join(lines)


def generate_markdown(job: TranscriptionJob) -> str:
    """Generate Markdown transcript."""
    lines = [
        f"# Transcript",
        f"",
        f"**Language:** {SUPPORTED_LANGUAGES.get(job.language, job.language)} ({job.language_probability*100:.1f}% confidence)",
        f"**Date:** {datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"",
        "---",
        ""
    ]

    current_speaker = None

    for segment in job.segments:
        speaker = segment.get("speaker")

        if speaker and speaker != current_speaker:
            lines.append(f"\n### {speaker}\n")
            current_speaker = speaker
        elif segment.get("paragraph_break"):
            lines.append("")

        timestamp = format_timestamp(segment["start"])
        lines.append(f"**[{timestamp}]** {segment['text']}\n")

    return "\n".join(lines)


def generate_srt(job: TranscriptionJob) -> str:
    """Generate SRT subtitle file."""
    lines = []

    for i, segment in enumerate(job.segments, 1):
        start = format_srt_timestamp(segment["start"])
        end = format_srt_timestamp(segment["end"])

        speaker_prefix = f"{segment['speaker']}: " if segment.get("speaker") else ""

        lines.append(str(i))
        lines.append(f"{start} --> {end}")
        lines.append(f"{speaker_prefix}{segment['text']}")
        lines.append("")

    return "\n".join(lines)


def format_vtt_timestamp(seconds: float) -> str:
    """Format seconds to WebVTT timestamp (HH:MM:SS.mmm with dot separator)."""
    hrs = int(seconds // 3600)
    mins = int((seconds % 3600) // 60)
    secs = int(seconds % 60)
    ms = int((seconds % 1) * 1000)
    return f"{hrs:02d}:{mins:02d}:{secs:02d}.{ms:03d}"


def generate_vtt(job: TranscriptionJob) -> str:
    """Generate WebVTT subtitle file."""
    lines = ["WEBVTT", ""]

    for i, segment in enumerate(job.segments, 1):
        start = format_vtt_timestamp(segment["start"])
        end = format_vtt_timestamp(segment["end"])

        speaker_prefix = f"{segment['speaker']}: " if segment.get("speaker") else ""

        lines.append(str(i))
        lines.append(f"{start} --> {end}")
        lines.append(f"{speaker_prefix}{segment['text']}")
        lines.append("")

    return "\n".join(lines)


def generate_pdf(job: TranscriptionJob) -> bytes:
    """Generate PDF transcript."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.75*inch, bottomMargin=0.75*inch)

    styles = getSampleStyleSheet()
    title_style = styles['Heading1']
    meta_style = ParagraphStyle('Meta', parent=styles['Normal'], fontSize=10, textColor='gray')
    speaker_style = ParagraphStyle('Speaker', parent=styles['Heading3'], fontSize=12, spaceAfter=6)
    text_style = ParagraphStyle('Text', parent=styles['Normal'], fontSize=11, leading=14, spaceAfter=12)
    timestamp_style = ParagraphStyle('Timestamp', parent=styles['Normal'], fontSize=9, textColor='blue')

    story = []

    story.append(Paragraph("Transcript", title_style))
    story.append(Spacer(1, 12))

    lang_name = SUPPORTED_LANGUAGES.get(job.language, job.language)
    story.append(Paragraph(f"Language: {lang_name} ({job.language_probability*100:.1f}% confidence)", meta_style))
    story.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", meta_style))
    story.append(Spacer(1, 24))

    current_speaker = None

    for segment in job.segments:
        speaker = segment.get("speaker")

        if segment.get("paragraph_break"):
            story.append(Spacer(1, 18))

        if speaker and speaker != current_speaker:
            story.append(Spacer(1, 12))
            story.append(Paragraph(speaker, speaker_style))
            current_speaker = speaker

        timestamp = format_timestamp(segment["start"])
        story.append(Paragraph(f"[{timestamp}]", timestamp_style))
        story.append(Paragraph(segment["text"], text_style))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def generate_docx(job: TranscriptionJob) -> bytes:
    """Generate DOCX transcript."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading("Transcript", level=1)

    lang_name = SUPPORTED_LANGUAGES.get(job.language, job.language)
    meta = doc.add_paragraph()
    meta.add_run(f"Language: {lang_name} ({job.language_probability*100:.1f}% confidence)\n").italic = True
    meta.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True

    doc.add_paragraph()

    current_speaker = None

    for segment in job.segments:
        speaker = segment.get("speaker")

        if segment.get("paragraph_break"):
            doc.add_paragraph()

        if speaker and speaker != current_speaker:
            doc.add_heading(speaker, level=2)
            current_speaker = speaker

        para = doc.add_paragraph()

        timestamp_run = para.add_run(f"[{format_timestamp(segment['start'])}] ")
        timestamp_run.font.color.rgb = RGBColor(0, 102, 204)
        timestamp_run.font.size = Pt(9)

        text_run = para.add_run(segment["text"])
        text_run.font.size = Pt(11)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def generate_json_export(
    job: TranscriptionJob,
    engine: str = None,
    warnings: list = None,
    edits: list = None,
) -> str:
    """Generate canonical JSON export with metadata and audit trail."""
    lang_name = SUPPORTED_LANGUAGES.get(job.language, job.language)

    data = {
        "metadata": {
            "generated_at": datetime.now().isoformat(),
            "engine": engine,
            "language": job.language,
            "language_name": lang_name,
            "language_confidence": job.language_probability,
            "segment_count": len(job.segments),
            "speaker_count": len(set(
                s.get("speaker") for s in job.segments if s.get("speaker")
            )),
        },
        "text": job.result,
        "segments": job.segments,
        "warnings": warnings or [],
        "edits": edits or [],
    }

    return json.dumps(data, indent=2, ensure_ascii=False)
